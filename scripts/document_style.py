"""Reusable, content-free DS-style typography for new manuals."""
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH


def font(style,latin='Arial',chinese='宋体',size=9.5,bold=False):
    style.font.name=latin; style.font.size=Pt(size); style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'),chinese)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rpr.rFonts.attrib.pop(qn('w:'+key),None)


def apply(doc,name):
    for s in doc.sections:
        s.top_margin=s.bottom_margin=Cm(2); s.left_margin=s.right_margin=Cm(2.5)
        s.header_distance=s.footer_distance=Cm(.9)
        s.different_first_page_header_footer=False
    for key,size in [('Normal',9.5),('Caption',9),('Title',18)]:
        st=doc.styles[key]; font(st,size=size,bold=(key=='Title'))
        st.paragraph_format.line_spacing=1.5
        st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(4)
    doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for key,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
        st=doc.styles[key]; font(st,chinese='微软雅黑',size=size,bold=True)
        pf=st.paragraph_format; pf.line_spacing=1.3; pf.space_before=Pt(12); pf.space_after=Pt(8)
        pf.keep_with_next=True
    for key in ('TOC 1','TOC 2','TOC 3','Table of Figures'):
        if key in doc.styles:
            st=doc.styles[key]; font(st,size=9)
            st.paragraph_format.line_spacing=1.4; st.paragraph_format.space_after=Pt(4)
    for p in doc.paragraphs:
        if p.style.name=='Caption':
            p.paragraph_format.line_spacing=1.4; p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(4)
            for r in p.runs: r.bold=False
    for t in doc.tables:
        widths=[c.width for c in t.columns]; total=sum(widths)
        fitted=[int(w*Cm(16)/total) for w in widths]
        for col,w in zip(t.columns,fitted): col.width=w
        from docx.enum.table import WD_TABLE_ALIGNMENT
        t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,row in enumerate(t.rows):
            for ci,c in enumerate(row.cells):
                c.width=fitted[ci]
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing=1.4
                    p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(1.5)
                    for r in p.runs:
                        r.font.name='Arial'; r.font.size=Pt(9); r.bold=(i==0)
                        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    for pic in doc.inline_shapes:
        if pic.width>Cm(16):
            pic.height=int(pic.height*Cm(16)/pic.width); pic.width=Cm(16)
    sec=doc.sections[0]
    header=sec.header.paragraphs[0]; header.text=name+' User Manual'; header.alignment=2
    footer=sec.footer.paragraphs[0]; footer.clear(); footer.alignment=1
    for p,edge in [(header,'bottom'),(footer,'top')]:
        p.paragraph_format.line_spacing=1; p.paragraph_format.space_after=Pt(2)
        b=OxmlElement('w:pBdr'); e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','6'),('space','3'),('color','000000')]: e.set(qn('w:'+k),v)
        b.append(e); p._p.get_or_add_pPr().append(b)
    for instr in ('PAGE','NUMPAGES'):
        if instr=='NUMPAGES': footer.add_run('/')
        f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),instr); footer._p.append(f)
    for p in (header,footer):
        for r in p.runs: r.font.name='Arial'; r.font.size=Pt(9)
