"""Create a new UM from a flat register model; never updates existing DOCX."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from figure_tools import render, register_figure
from document_style import apply as apply_document_style


ACCESS={'read-only','write-only','read-write','writeOnce','read-writeOnce'}

def validate(model):
    seen=set(); spans=[]
    for p in model['peripherals']:
        if p['name'] in seen: raise ValueError('Duplicate peripheral')
        seen.add(p['name']); names=set()
        for r in p['registers']:
            if r['name'] in names: raise ValueError('Duplicate register')
            names.add(r['name'])
            for k in ('offset','size','address','reset_value','reset_mask'):
                if type(r[k]) is not int: raise ValueError('Register integer required: '+k)
            n=r['size']; addr=p['base_address']+r['offset']
            if n not in (8,16,32) or r['offset']<0 or addr!=r['address'] or addr%(n//8): raise ValueError('Address/width mismatch')
            if r['access'] not in ACCESS: raise ValueError('Unsupported access')
            if not (0<=r['reset_value']<2**n and 0<=r['reset_mask']<2**n): raise ValueError('Reset overflow')
            for a,b in spans:
                if addr<b and addr+n//8>a: raise ValueError('Overlapping register addresses')
            spans.append((addr,addr+n//8)); bits=0; fields=set()
            for f in r.get('fields',[]):
                l,w=f['lsb'],f['width']
                if type(l) is not int or type(w) is not int or l<0 or w<1 or l+w>n: raise ValueError('Field range')
                mask=((1<<w)-1)<<l
                if bits&mask or f['name'] in fields: raise ValueError('Overlapping/duplicate fields')
                bits|=mask; fields.add(f['name']); enums=set()
                if f.get('access') and f['access'] not in ACCESS: raise ValueError('Field access')
                for e in f.get('enums',[]):
                    if type(e['value']) is not int or not 0<=e['value']<2**w or e['name'] in enums: raise ValueError('Invalid enum')
                    enums.add(e['name'])


def field(p, code, cached=''):
    e=OxmlElement('w:fldSimple'); e.set(qn('w:instr'),code)
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=str(cached); r.append(t); e.append(r); p._p.append(e)


def bookmark(p,name,index):
    a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),str(index)); a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'),str(index)); p._p.insert(0,a); p._p.append(b)


def build(model_path,content_path,out):
    model_path=Path(model_path); content_path=Path(content_path); out=Path(out)
    model=json.loads(model_path.read_text(encoding='utf-8-sig')); content=json.loads(content_path.read_text(encoding='utf-8-sig'))
    validate(model)
    if type(content.get('synthetic')) is not bool: raise ValueError('Explicit synthetic flag required')
    if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_-]*',model['device']['name']): raise ValueError('Unsafe device filename')
    # Check all figure references before creating output.
    figures=[]; ids=set()
    for f in content.get('figures',[]):
        if not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]*',f['id']) or f['id'] in ids: raise ValueError('Invalid/duplicate figure id')
        ids.add(f['id']); src=(content_path.parent/f['diagram']).resolve()
        if not src.is_relative_to(content_path.parent.resolve()) or not src.is_file(): raise ValueError('Figure path outside input folder or missing')
        figures.append((f,src))
    out.mkdir(parents=True,exist_ok=False)
    doc=Document(); sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
    style=content.get('style',{}); body=style.get('body_font','宋体'); latin=style.get('latin_font','Times New Roman')
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
        st=doc.styles[name]; st.font.name=latin; st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),body)
        for border in st.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    normal=doc.styles['Normal']; normal.font.size=Pt(style.get('body_size',10.5)); normal.paragraph_format.line_spacing=style.get('line_spacing',1.25)
    normal.paragraph_format.space_after=Pt(5)
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''; doc.core_properties.title=model['device']['name']+' 用户手册'
    captions={'图':0,'表':0}; gaps=[]
    def caption(label,title):
        captions[label]+=1; num=captions[label]; p=doc.add_paragraph(style='Caption'); p.alignment=1
        p.paragraph_format.keep_with_next=(label=='表')
        p.add_run(label+' '); field(p,'SEQ '+label+' \\* ARABIC',num); p.add_run(' '+title)
        bookmark(p,('fig_' if label=='图' else 'tbl_')+str(num),1000+(0 if label=='图' else 500)+num)
        return num
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
        if widths:
            for c,w in zip(t.columns,widths): c.width=Cm(w)
        for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
        for row in rows:
            for c,v in zip(t.add_row().cells,row): c.text=str(v)
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),style.get('table_border_color','000000')); borders.append(e)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
            if i==0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for ci,c in enumerate(row.cells):
                if widths: c.width=Cm(widths[ci])
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i==0:
                    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),style.get('table_header_fill','E0E0E0')); c._tc.get_or_add_tcPr().append(sh)
                for p in c.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                    for run in p.runs: run.font.size=Pt(9); run.bold=(i==0)
        return t
    def picture(path,title):
        p=doc.add_paragraph(); p.alignment=1; p.paragraph_format.keep_with_next=True
        shape=p.add_run().add_picture(str(path),width=Cm(16.4)); shape._inline.docPr.set('descr',title)
        sp=shape._inline.xpath('.//pic:spPr')[0]; line=OxmlElement('a:ln'); line.set('w','6350'); fill=OxmlElement('a:solidFill'); clr=OxmlElement('a:srgbClr'); clr.set('val','000000'); fill.append(clr); line.append(fill); sp.append(line)
        return caption('图',title)
    doc.add_paragraph(model['device']['name']+' 用户手册',style='Title')
    doc.add_paragraph('版本 '+model['device'].get('version','未提供'))
    doc.add_paragraph(content['scope'])
    if content['synthetic']: doc.add_paragraph('本手册使用虚构芯片和合成寄存器数据，仅用于展示文档制作流程，不对应真实芯片。')
    doc.add_page_break()
    doc.add_paragraph('目录',style='Title'); field(doc.add_paragraph(),'TOC \\o "1-3" \\h \\z \\u','请更新目录域')
    doc.add_page_break()
    doc.add_paragraph('插图目录',style='Title'); field(doc.add_paragraph(),'TOC \\h \\z \\c "图"','请更新插图目录域')
    doc.add_paragraph('表格目录',style='Title'); field(doc.add_paragraph(),'TOC \\h \\z \\c "表"','请更新表格目录域')
    doc.add_page_break(); doc.add_heading('1 系统概述',level=1); doc.add_paragraph(model['device']['description'])
    sources=[model_path,content_path]
    for f,src in figures:
        render(json.loads(src.read_text(encoding='utf-8-sig')),out/'figures'/f['id']); sources.append(src)
        num=picture(out/'figures'/(f['id']+'.png'),f['title'])
        p=doc.add_paragraph('系统结构见'); field(p,f'REF fig_{num} \\h',f'图 {num} {f["title"]}')
    for pi,p in enumerate(model['peripherals'],2):
        doc.add_heading(f'{pi} {p["name"]}',level=1)
        c=content.get('peripherals',{}).get(p['name'],{})
        doc.add_paragraph(c.get('overview',p.get('description','')))
        if not c: gaps.append(p['name']+': missing functional description')
        if not p.get('source'): gaps.append(p['name']+': missing source')
        for feat in c.get('features',[]): doc.add_paragraph(feat,style='List Bullet')
        if c.get('steps'):
            doc.add_heading('配置流程',level=2)
            for j,step in enumerate(c['steps'],1): doc.add_paragraph(f'{j}. {step}')
        caption('表',p['name']+' 寄存器映射')
        table(['寄存器','偏移','绝对地址','访问','复位值'],[[r['name'],f'0x{r["offset"]:X}',f'0x{r["address"]:08X}',r['access'],f'0x{r["reset_value"]:08X}'] for r in p['registers']])
        for ri,r in enumerate(p['registers']):
            doc.add_heading(r['name']+' 寄存器',level=2)
            doc.add_paragraph(r.get('description',''))
            doc.add_paragraph(f'复位有效掩码 0x{r["reset_mask"]:08X}。掩码外位的复位状态未知；未列出的位不据此认定为保留位。')
            stem=out/'figures'/f'reg_{pi}_{ri}'; register_figure(r,stem); picture(stem.with_suffix('.png'),p['name']+' '+r['name']+' 位分布')
            rows=[]
            for f in r.get('fields',[]):
                desc=f.get('description','')
                access=f.get('access') or r['access']
                for key,label in [('modified_write_values','写副作用'),('read_action','读副作用')]:
                    val=f.get(key) or r.get(key) or '未声明'
                    desc+='\n'+label+'：'+val
                for e in f.get('enums',[]): desc+=f'\n{e["value"]} {e["name"]}：{e["description"]}'
                mask=((1<<f['width'])-1)<<f['lsb']
                reset=str((r['reset_value']&mask)>>f['lsb']) if r['reset_mask']&mask==mask else '部分或全部未知'
                rows.append([f'{f["lsb"]+f["width"]-1}:{f["lsb"]}',f['name'],access,reset,desc])
            caption('表',p['name']+' '+r['name']+' 位域说明'); table(['位','字段','访问','复位','说明'],rows,[1.1,2.2,2.5,1.8,9])
    doc.add_heading('来源与版本记录',level=1)
    doc.add_paragraph('寄存器和功能说明分别维护。下面的来源标识用于追溯示例定义。')
    for p in model['peripherals']: doc.add_paragraph(p['name']+'：'+p.get('source','来源待补'))
    settings=doc.settings.element; flag=OxmlElement('w:updateFields'); flag.set(qn('w:val'),'true'); settings.append(flag)
    footer=sec.footer.paragraphs[0]; footer.alignment=1; field(footer,'PAGE','1')
    if style.get('profile','ds')=='ds': apply_document_style(doc,model['device']['name'])
    dest=out/(model['device']['name']+'_UM.docx'); doc.save(dest)
    report={'model_valid':True,'gaps':gaps,'fields_updated':False,'visual_review_completed':False}
    (out/'build_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    manifest={'inputs':[{'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sources], 'outputs':[{'file':p.relative_to(out).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(out.rglob('*')) if p.is_file()]}
    (out/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    return dest


if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--model',required=True); ap.add_argument('--content',required=True); ap.add_argument('--out',required=True)
    a=ap.parse_args(); print(build(a.model,a.content,a.out))
